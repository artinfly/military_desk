"""Массовые операции воинского учёта (как в старой программе, меню «Операции»).

Каждая операция выполняется целиком в одной транзакции: или всё, или ничего.
"""

import re
from dataclasses import dataclass, field

from django.db import transaction
from openpyxl import load_workbook

from apps.references.models import SpecialAccountType

from .models import MilitaryRecord, ReservationDismissal

# Код группы «Общий учёт» — её ставят при разбронировании
GENERAL_ACCOUNT_CODE = "H9"


def snils_digits(value):
    """СНИЛС без пробелов и дефисов: «926-581-417 72» → «92658141772»."""
    return re.sub(r"\D", "", str(value or ""))


@transaction.atomic
def change_account_type(old_type, new_type):
    """Смена группы учёта у незабронированных: old_type → new_type (None — убрать группу).

    Возвращает число изменённых записей.
    """
    return MilitaryRecord.objects.filter(
        special_account_type=old_type, is_reserved=False
    ).update(special_account_type=new_type)


@dataclass
class DismissedUpdate:
    removed: int  # снова работают — убраны из спец. уволенных
    added: int  # уволенные забронированные — добавлены
    unreserved: int  # уволенным снята бронь


@transaction.atomic
def update_reserved_dismissed():
    """Обновление списка спец. уволенных (увольнений забронированных работников).

    1. Кто снова работает — убрать из списка.
    2. Уволенных забронированных — добавить (удостоверение и дата увольнения),
       если такой записи ещё нет.
    3. У всех уволенных снять бронь.
    """
    removed, _ = ReservationDismissal.objects.filter(
        person__dismissal_date__isnull=True
    ).delete()

    reserved = MilitaryRecord.objects.filter(
        is_reserved=True, person__dismissal_date__isnull=False
    ).select_related("person")
    existing = set(ReservationDismissal.objects.values_list("person_id", "order"))
    new = [
        ReservationDismissal(
            person=record.person,
            order=record.reservation_document,
            dismissal_date=record.person.dismissal_date,
        )
        for record in reserved
        if (record.person_id, record.reservation_document) not in existing
    ]
    ReservationDismissal.objects.bulk_create(new)

    unreserved = MilitaryRecord.objects.filter(
        is_reserved=True, person__dismissal_date__isnull=False
    ).update(is_reserved=False)
    return DismissedUpdate(removed, len(new), unreserved)


def read_snils_file(file):
    """СНИЛС из файла Excel: первая колонка, со второй строки (первая — заголовок)."""
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        values = (
            row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)
        )
        return [snils for snils in map(snils_digits, values) if snils]
    finally:
        book.close()


@dataclass
class UnreserveResult:
    unreserved: int = 0
    not_found: list = field(default_factory=list)  # СНИЛС, которых нет среди работающих


@transaction.atomic
def unreserve_by_snils(snils_list, order):
    """Разбронирование работающих по списку СНИЛС.

    Бронь снимается, группа учёта — «Общий учёт», в карточку пишется номер
    приказа о разбронировании, удостоверение об отсрочке очищается.
    """
    general = SpecialAccountType.objects.filter(code=GENERAL_ACCOUNT_CODE).first()
    records = MilitaryRecord.objects.filter(
        person__dismissal_date__isnull=True
    ).exclude(person__snils="")
    by_snils = {}
    for pk, snils in records.values_list("pk", "person__snils"):
        by_snils.setdefault(snils_digits(snils), []).append(pk)

    result = UnreserveResult()
    found = []
    for snils in dict.fromkeys(snils_list):  # без повторов, порядок файла
        if snils in by_snils:
            found.extend(by_snils[snils])
        else:
            result.not_found.append(snils)
    result.unreserved = MilitaryRecord.objects.filter(pk__in=found).update(
        is_reserved=False,
        special_account_type=general,
        unreservation_order=order,
        reservation_document="",
    )
    return result
