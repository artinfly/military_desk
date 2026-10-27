"""Общая сборка отчётов Excel: оформленные таблицы и ответ для скачивания."""

from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2E2E2E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _fill_sheet(sheet, title, headers, rows):
    """Лист: шапка (закреплена, с автофильтром), строки, ширина колонок по данным."""
    sheet.title = title[:31]  # ограничение Excel на имя листа
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # Ширина колонки — по самому длинному значению (в разумных пределах)
    for index, header in enumerate(headers, 1):
        values = [len(str(header))] + [
            len(str(r[index - 1])) for r in rows if r[index - 1] is not None
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(values) + 2, 60
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if hasattr(cell.value, "year"):
                cell.number_format = "DD.MM.YYYY"


def build_sheets(sheets):
    """Книга из листов [(название, шапка, строки)]. Возвращает байты файла."""
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sheets:
        _fill_sheet(book.create_sheet(), title, headers, rows)
    content = BytesIO()
    book.save(content)
    return content.getvalue()


def build_workbook(title, headers, rows):
    """Книга с одним листом."""
    return build_sheets([(title, headers, rows)])


def xlsx_response(content, filename):
    """Ответ «скачать файл»; имя файла может быть по-русски."""
    response = HttpResponse(content, content_type=XLSX)
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response
